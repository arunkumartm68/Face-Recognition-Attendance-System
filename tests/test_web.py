"""Tests for the web app, database, reports, exports and bulk import.

Run from the project folder:  python -m unittest discover -s tests -v
"""
import _env  # noqa: F401  - must be imported before the app

import io
import json
import os
import sqlite3
import tempfile
import time
import unittest
from datetime import datetime, timedelta
from unittest import mock

from openpyxl import load_workbook

import app as webapp
import common
import reports

TODAY = datetime.now().strftime(common.DATE_FORMAT)
PASSWORD = 'secret1'
PASSWORD_HASH = common.hash_password(PASSWORD)
TOKEN = 'test-csrf-token'


def days_ago(n):
    return (datetime.now() - timedelta(days=n)).strftime(common.DATE_FORMAT)


class FakeProcess:
    """Stands in for subprocess.Popen so no window or camera is opened."""

    def __init__(self, args, cwd=None):
        self.args, self.cwd, self.running = args, cwd, True

    def poll(self):
        return None if self.running else 0


class WebTestCase(unittest.TestCase):
    def setUp(self):
        _env.reset_state()
        webapp._processes.clear()
        webapp._process_info.clear()
        self.started = []

        def fake_popen(args, cwd=None):
            proc = FakeProcess(args, cwd)
            self.started.append(proc)
            return proc

        patcher = mock.patch.object(webapp.subprocess, 'Popen', side_effect=fake_popen)
        patcher.start()
        self.addCleanup(patcher.stop)
        models = mock.patch.object(common, 'missing_model_files', return_value=[])
        models.start()
        self.addCleanup(models.stop)

    # -- data helpers (direct SQL keeps the tests fast) --------------------

    def db(self):
        return common.connect()

    def add_student(self, reg, name):
        with self.db() as conn:
            conn.execute('INSERT INTO students VALUES (?, ?, ?)', (reg, name, PASSWORD_HASH))

    def add_subject(self, subject, professor='Prof'):
        with self.db() as conn:
            conn.execute('INSERT INTO staff VALUES (?, ?, ?)', (subject, professor, PASSWORD_HASH))

    def mark(self, reg, subject, date, status='Present', time_='09:00:00 AM'):
        with self.db() as conn:
            conn.execute('INSERT INTO attendance (register_number, name, time, date, subject, status) '
                         'VALUES (?, ?, ?, ?, ?, ?)', (reg, reg, time_, date, subject, status))

    def enroll(self, reg, subject):
        with self.db() as conn:
            conn.execute('INSERT INTO enrollments VALUES (?, ?)', (reg, subject))

    def query(self, sql, args=()):
        conn = self.db()
        try:
            return conn.execute(sql, args).fetchall()
        finally:
            conn.close()

    # -- client helpers ------------------------------------------------------

    def login(self, role, username, password=PASSWORD):
        client = webapp.app.test_client()
        with client.session_transaction() as s:
            s['_csrf_token'] = TOKEN
        response = client.post('/login', data={'role': role, 'username': username,
                                               'password': password, 'csrf_token': TOKEN})
        return client, response

    def admin(self):
        return self.login('admin', 'admin', 'admin')[0]

    @staticmethod
    def post(client, url, data=None, **kwargs):
        with client.session_transaction() as s:
            token = s.setdefault('_csrf_token', TOKEN)
        return client.post(url, data=dict(data or {}, csrf_token=token), **kwargs)

    @staticmethod
    def post_json(client, url, payload=None):
        with client.session_transaction() as s:
            token = s.setdefault('_csrf_token', TOKEN)
        return client.post(url, json=payload or {},
                           headers={'X-CSRF-Token': token, 'Accept': 'application/json'})


class TestDatabase(WebTestCase):
    def test_init_creates_tables_and_hashed_default_admin(self):
        tables = {r[0] for r in self.query("SELECT name FROM sqlite_master WHERE type = 'table'")}
        self.assertTrue({'students', 'staff', 'admins', 'attendance', 'enrollments',
                         'sessions', 'audit_log'} <= tables)
        admin = self.query('SELECT * FROM admins')
        self.assertEqual(len(admin), 1)
        self.assertTrue(common.is_password_hash(admin[0]['password']))
        self.assertTrue(common.verify_password(admin[0]['password'], 'admin'))

    def test_init_is_idempotent(self):
        common.init_db()
        common.init_db()
        self.assertEqual(self.query('SELECT COUNT(*) FROM admins')[0][0], 1)

    def test_migrates_legacy_database(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'legacy.db')
            conn = sqlite3.connect(path)
            conn.executescript('''
                CREATE TABLE students (register_number TEXT PRIMARY KEY, name TEXT, password TEXT);
                CREATE TABLE staff (subject_name TEXT PRIMARY KEY, password TEXT);
                CREATE TABLE attendance (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, time TEXT,
                                         date TEXT, subject TEXT, UNIQUE(name, date, subject));
                INSERT INTO students VALUES ('R1', 'Arun', 'plain1');
                INSERT INTO staff VALUES ('Maths', 'plain2');
                INSERT INTO attendance (name, time, date, subject) VALUES ('R1_Arun', '09:00:00 AM', '2026-01-05', 'Maths');
            ''')
            conn.commit()
            conn.close()

            common.init_db(path)

            conn = common.connect(path)
            try:
                student = conn.execute('SELECT * FROM students').fetchone()
                self.assertTrue(common.is_password_hash(student['password']))
                self.assertTrue(common.verify_password(student['password'], 'plain1'))
                staff = conn.execute('SELECT * FROM staff').fetchone()
                self.assertEqual(staff['professor_name'], '')
                self.assertTrue(common.verify_password(staff['password'], 'plain2'))
                row = conn.execute('SELECT * FROM attendance').fetchone()
                self.assertEqual((row['register_number'], row['name'], row['status']), ('R1', 'Arun', 'Present'))
            finally:
                conn.close()

    def test_adds_status_column_to_previous_schema(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'v1.db')
            conn = sqlite3.connect(path)
            conn.executescript('''
                CREATE TABLE attendance (id INTEGER PRIMARY KEY AUTOINCREMENT, register_number TEXT NOT NULL,
                    name TEXT, time TEXT, date TEXT, subject TEXT DEFAULT 'General',
                    UNIQUE(register_number, date, subject));
                INSERT INTO attendance (register_number, name, time, date, subject)
                    VALUES ('R1', 'Arun', '09:00:00 AM', '2026-01-05', 'General');
            ''')
            conn.commit()
            conn.close()
            common.init_db(path)
            conn = common.connect(path)
            try:
                self.assertEqual(conn.execute('SELECT status FROM attendance').fetchone()[0], 'Present')
            finally:
                conn.close()


class TestHelpers(unittest.TestCase):
    def test_validation(self):
        self.assertIsNone(common.validate_register_number('21CS001'))
        self.assertIn("'_'", common.validate_register_number('21_CS_001'))
        self.assertIsNotNone(common.validate_register_number('21/CS/001'))
        self.assertIsNotNone(common.validate_register_number(''))
        self.assertIsNone(common.validate_name("D'Souza"))
        self.assertIsNone(common.validate_name('José Kumar'))
        self.assertIsNotNone(common.validate_name('Ar:un'))
        self.assertIsNotNone(common.validate_password('12345'))
        self.assertIsNone(common.validate_password('123456'))
        self.assertIsNotNone(common.validate_subject_name('general'))

    def test_labels(self):
        self.assertEqual(common.folder_label('person_3_21CS001_Arun Kumar'), '21CS001_Arun Kumar')
        self.assertEqual(common.split_label('21CS001_Arun_K'), ('21CS001', 'Arun_K'))
        self.assertEqual(common.split_label('NoUnderscore'), ('NoUnderscore', 'NoUnderscore'))

    def test_face_folder_helpers(self):
        _env.reset_state()
        for folder in ('person_1_R1_Arun', 'person_2_R2_Priya', 'person_3_R1_Arun'):
            os.makedirs(os.path.join(common.FACES_DIR, folder))
        for i in range(3):
            open(os.path.join(common.FACES_DIR, 'person_1_R1_Arun', f'img_face_{i}.jpg'), 'wb').close()
        open(os.path.join(common.FACES_DIR, 'person_3_R1_Arun', 'img_face_1.jpg'), 'wb').close()
        self.assertEqual(common.face_photo_counts(), {'R1': 4, 'R2': 0})
        self.assertEqual(common.delete_face_folders('R1'), 2)
        self.assertEqual(sorted(os.listdir(common.FACES_DIR)), ['person_2_R2_Priya'])


class TestLoginAndSecurity(WebTestCase):
    def test_admin_default_password_shows_warning(self):
        client, response = self.login('admin', 'admin', 'admin')
        self.assertEqual(response.headers['Location'], '/admin-dashboard')
        self.assertIn('default admin password', client.get('/admin-dashboard').get_data(as_text=True))

    def test_wrong_password_is_rejected(self):
        client, response = self.login('admin', 'admin', 'wrong')
        self.assertEqual(response.headers['Location'], '/')
        self.assertIn('Invalid admin credentials', client.get('/').get_data(as_text=True))
        self.assertEqual(client.get('/admin-dashboard').status_code, 302)

    def test_staff_and_student_logins(self):
        self.add_subject('Maths')
        self.add_student('R1', 'Arun')
        self.assertEqual(self.login('staff', 'Maths')[1].headers['Location'], '/staff-dashboard')
        self.assertEqual(self.login('student', 'R1')[1].headers['Location'], '/student-dashboard')
        self.assertEqual(self.login('student', 'Arun')[1].headers['Location'], '/student-dashboard')
        self.assertEqual(self.login('student', 'R1', 'wrong')[1].headers['Location'], '/')

    def test_post_without_csrf_token_is_rejected(self):
        client = self.admin()
        response = client.post('/manage-student', data={'action': 'add', 'register_number': 'R1',
                                                         'name': 'Arun', 'password': PASSWORD})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.query('SELECT COUNT(*) FROM students')[0][0], 0)
        response = client.post('/run-register', headers={'Accept': 'application/json'})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()['status'], 'error')
        self.assertEqual(self.started, [])

    def test_roles_are_enforced(self):
        anonymous = webapp.app.test_client()
        self.assertEqual(anonymous.get('/admin-dashboard').status_code, 302)
        self.assertEqual(anonymous.get('/export/xlsx').status_code, 302)
        with anonymous.session_transaction() as s:
            s['_csrf_token'] = TOKEN
        response = anonymous.post('/run-attendance', json={}, headers={'X-CSRF-Token': TOKEN,
                                                                      'Accept': 'application/json'})
        self.assertEqual(response.status_code, 401)
        self.add_student('R1', 'Arun')
        student, _ = self.login('student', 'R1')
        self.assertEqual(student.get('/admin/students').status_code, 302)
        self.assertEqual(student.get('/staff-dashboard').status_code, 302)

    def test_deleted_student_is_logged_out(self):
        self.add_student('R1', 'Arun')
        student, _ = self.login('student', 'R1')
        self.assertEqual(student.get('/student-dashboard').status_code, 200)
        with self.db() as conn:
            conn.execute("DELETE FROM students WHERE register_number = 'R1'")
        response = student.get('/student-dashboard')
        self.assertEqual(response.headers['Location'], '/')

    def test_no_passwords_or_hashes_in_pages(self):
        self.add_student('R1', 'Arun')
        self.add_subject('Maths')
        client = self.admin()
        for page in ('/admin/students', '/admin/subjects', '/admin-dashboard'):
            html = client.get(page).get_data(as_text=True)
            self.assertNotIn('scrypt:', html)
            self.assertNotIn(PASSWORD, html)

    def test_security_headers_and_secret_key(self):
        response = webapp.app.test_client().get('/')
        self.assertEqual(response.headers['X-Frame-Options'], 'DENY')
        self.assertEqual(response.headers['X-Content-Type-Options'], 'nosniff')
        self.assertNotEqual(webapp.app.secret_key, 'super_secret_key')
        with open(common.SECRET_KEY_FILE, encoding='utf-8') as f:
            self.assertEqual(f.read().strip(), webapp.app.secret_key)

    def test_registration_latency_api(self):
        with open(common.LATENCY_JSON, 'w', encoding='utf-8') as f:
            json.dump({'registered': True, 'photos_count': 2, 'latencies': [40.1, 42.0],
                       'student_folder': 'R1_Arun'}, f)
        self.assertFalse(webapp.app.test_client().get('/api/registration-latency').get_json()['registered'])
        client = self.admin()
        data = client.get('/api/registration-latency').get_json()
        self.assertTrue(data['registered'])
        self.assertEqual(data['photos_target'], common.PHOTOS_PER_STUDENT)
        old = time.time() - 3600
        os.utime(common.LATENCY_JSON, (old, old))
        self.assertFalse(client.get('/api/registration-latency').get_json()['registered'])


class TestChangePassword(WebTestCase):
    def test_change_password_flow(self):
        client = self.admin()
        cases = [({'current_password': 'nope', 'new_password': 'abcdef1', 'confirm_password': 'abcdef1'}, 'incorrect'),
                 ({'current_password': 'admin', 'new_password': 'abcdef1', 'confirm_password': 'abcdef2'}, 'do not match'),
                 ({'current_password': 'admin', 'new_password': 'abc', 'confirm_password': 'abc'}, 'at least')]
        for data, message in cases:
            self.post(client, '/change-password', data)
            self.assertIn(message, client.get('/change-password').get_data(as_text=True))
        response = self.post(client, '/change-password', {'current_password': 'admin', 'new_password': 'abcdef1',
                                                          'confirm_password': 'abcdef1'})
        self.assertEqual(response.headers['Location'], '/admin-dashboard')
        self.assertNotIn('default admin password', client.get('/admin-dashboard').get_data(as_text=True))
        self.assertEqual(self.login('admin', 'admin', 'admin')[1].headers['Location'], '/')
        self.assertEqual(self.login('admin', 'admin', 'abcdef1')[1].headers['Location'], '/admin-dashboard')

    def test_student_changes_own_password(self):
        self.add_student('R1', 'Arun')
        client, _ = self.login('student', 'R1')
        self.post(client, '/change-password', {'current_password': PASSWORD, 'new_password': 'newpass1',
                                               'confirm_password': 'newpass1'})
        self.assertEqual(self.login('student', 'R1', 'newpass1')[1].headers['Location'], '/student-dashboard')


class TestStudents(WebTestCase):
    def test_add_validates_input(self):
        client = self.admin()
        bad = [('21_CS_001', 'Arun', PASSWORD), ('21/CS', 'Arun', PASSWORD), ('R1', 'Arun', '123')]
        for reg, name, password in bad:
            self.post(client, '/manage-student', {'action': 'add', 'register_number': reg, 'name': name,
                                                  'password': password})
        self.assertEqual(self.query('SELECT COUNT(*) FROM students')[0][0], 0)
        self.post(client, '/manage-student', {'action': 'add', 'register_number': 'R1', 'name': 'Arun',
                                              'password': PASSWORD})
        self.post(client, '/manage-student', {'action': 'add', 'register_number': 'R1', 'name': 'Other',
                                              'password': PASSWORD})
        rows = self.query('SELECT * FROM students')
        self.assertEqual([(r['register_number'], r['name']) for r in rows], [('R1', 'Arun')])
        self.assertTrue(common.verify_password(rows[0]['password'], PASSWORD))

    def test_edit_keeps_or_changes_password(self):
        self.add_student('R1', 'Arun')
        self.mark('R1', 'General', TODAY)
        client = self.admin()
        self.post(client, '/manage-student', {'action': 'edit', 'register_number': 'R1', 'name': 'Arun K',
                                              'password': ''})
        row = self.query('SELECT * FROM students')[0]
        self.assertEqual(row['name'], 'Arun K')
        self.assertTrue(common.verify_password(row['password'], PASSWORD))
        self.assertEqual(self.query('SELECT name FROM attendance')[0][0], 'Arun K')
        self.post(client, '/manage-student', {'action': 'edit', 'register_number': 'R1', 'name': 'Arun K',
                                              'password': 'changed1'})
        self.assertTrue(common.verify_password(self.query('SELECT password FROM students')[0][0], 'changed1'))

    def test_delete_removes_history_photos_and_rebuilds_features(self):
        self.add_subject('Maths')
        self.add_student('R1', 'Arun')
        self.add_student('R2', 'Priya')
        self.mark('R1', 'Maths', TODAY)
        self.mark('R2', 'Maths', TODAY)
        self.enroll('R1', 'Maths')
        os.makedirs(os.path.join(common.FACES_DIR, 'person_1_R1_Arun'))
        os.makedirs(os.path.join(common.FACES_DIR, 'person_2_R2_Priya'))
        client = self.admin()
        self.post(client, '/manage-student', {'action': 'delete', 'register_number': 'R1'})
        self.assertEqual([r[0] for r in self.query('SELECT register_number FROM students')], ['R2'])
        self.assertEqual([r[0] for r in self.query('SELECT register_number FROM attendance')], ['R2'])
        self.assertEqual(self.query('SELECT COUNT(*) FROM enrollments')[0][0], 0)
        self.assertEqual(os.listdir(common.FACES_DIR), ['person_2_R2_Priya'])
        self.assertEqual(self.started[-1].args[1], 'features_extraction_to_csv.py')
        log = self.query('SELECT * FROM audit_log')[0]
        self.assertEqual((log['action'], log['register_number']), ('Deleted student', 'R1'))
        self.assertIn('1 attendance record', log['details'])

    def test_students_page_shows_photo_status(self):
        self.add_student('R1', 'Arun')
        folder = os.path.join(common.FACES_DIR, 'person_1_R1_Arun')
        os.makedirs(folder)
        for i in range(common.PHOTOS_PER_STUDENT):
            open(os.path.join(folder, f'img_face_{i}.jpg'), 'wb').close()
        html = self.admin().get('/admin/students').get_data(as_text=True)
        self.assertIn(f'{common.PHOTOS_PER_STUDENT} photos', html)


class TestSubjects(WebTestCase):
    def test_add_edit_and_reserved_name(self):
        client = self.admin()
        self.post(client, '/manage-staff', {'action': 'add', 'subject_name': 'General', 'password': PASSWORD})
        self.post(client, '/manage-staff', {'action': 'add', 'subject_name': 'Maths', 'professor_name': 'Ravi',
                                            'password': PASSWORD})
        self.post(client, '/manage-staff', {'action': 'edit', 'subject_name': 'Maths', 'professor_name': 'Ravi K',
                                            'password': ''})
        rows = self.query('SELECT * FROM staff')
        self.assertEqual([(r['subject_name'], r['professor_name']) for r in rows], [('Maths', 'Ravi K')])
        self.assertTrue(common.verify_password(rows[0]['password'], PASSWORD))

    def test_delete_removes_subject_data(self):
        self.add_subject('Maths')
        self.add_subject('Physics')
        self.add_student('R1', 'Arun')
        self.mark('R1', 'Maths', TODAY)
        self.mark('R1', 'Physics', TODAY)
        self.enroll('R1', 'Maths')
        with self.db() as conn:
            conn.execute("INSERT INTO sessions (subject, date, started_at) VALUES ('Maths', ?, ?)",
                         (TODAY, TODAY + ' 09:00:00'))
        self.post(self.admin(), '/manage-staff', {'action': 'delete', 'subject_name': 'Maths'})
        self.assertEqual([r[0] for r in self.query('SELECT subject_name FROM staff')], ['Physics'])
        self.assertEqual([r[0] for r in self.query('SELECT subject FROM attendance')], ['Physics'])
        self.assertEqual(self.query('SELECT COUNT(*) FROM sessions')[0][0], 0)
        self.assertEqual(self.query('SELECT COUNT(*) FROM enrollments')[0][0], 0)
        self.assertEqual(self.query('SELECT action FROM audit_log')[0][0], 'Deleted subject')

    def test_enrollment_controls_the_roster(self):
        self.add_subject('Maths')
        for reg in ('R1', 'R2', 'R3'):
            self.add_student(reg, reg)
        conn = self.db()
        try:
            self.assertEqual(len(common.roster(conn, 'Maths')), 3)   # nobody assigned -> everyone
        finally:
            conn.close()
        client = self.admin()
        with client.session_transaction() as s:
            s['_csrf_token'] = TOKEN
        client.post('/admin/subjects/enrollment', data={'csrf_token': TOKEN, 'subject_name': 'Maths',
                                                        'register_numbers': ['R1', 'R3', 'ghost']})
        conn = self.db()
        try:
            self.assertEqual([r[0] for r in common.roster(conn, 'Maths')], ['R1', 'R3'])
            self.assertTrue(common.on_roster(conn, 'R1', 'Maths'))
            self.assertFalse(common.on_roster(conn, 'R2', 'Maths'))
            self.assertTrue(common.on_roster(conn, 'R2', common.GENERAL_SUBJECT))
            self.assertEqual(len(common.roster(conn, common.GENERAL_SUBJECT)), 3)
        finally:
            conn.close()
        html = client.get('/admin-dashboard?subject=Maths&date=' + TODAY).get_data(as_text=True)
        self.assertIn('R3', html)
        self.assertNotIn('<td>R2</td>', html)
        client.post('/admin/subjects/enrollment', data={'csrf_token': TOKEN, 'subject_name': 'Maths'})
        self.assertEqual(self.query('SELECT COUNT(*) FROM enrollments')[0][0], 0)


class TestAttendanceChanges(WebTestCase):
    def setUp(self):
        super().setUp()
        self.add_subject('Maths')
        self.add_student('R1', 'Arun')
        self.add_student('R2', 'Priya')

    def status(self, reg, subject='Maths', date=TODAY):
        rows = self.query('SELECT status, time FROM attendance WHERE register_number = ? AND subject = ? AND date = ?',
                          (reg, subject, date))
        return rows[0] if rows else None

    def test_admin_marks_any_date_and_changes_are_logged(self):
        client = self.admin()
        date = days_ago(3)
        self.post(client, '/admin/attendance', {'register_number': 'R1', 'subject': 'Maths', 'date': date,
                                                'status': 'Present'})
        first_time = self.status('R1', date=date)['time']
        self.post(client, '/admin/attendance', {'register_number': 'R1', 'subject': 'Maths', 'date': date,
                                                'status': 'Late'})
        row = self.status('R1', date=date)
        self.assertEqual((row['status'], row['time']), ('Late', first_time))
        self.post(client, '/admin/attendance', {'register_number': 'R1', 'subject': 'Maths', 'date': date,
                                                'status': 'Absent'})
        self.assertIsNone(self.status('R1', date=date))
        log = self.query('SELECT action, details, actor_role FROM audit_log ORDER BY id')
        self.assertEqual([(r['action'], r['details']) for r in log],
                         [('Marked Present', 'was Absent'), ('Marked Late', 'was Present'),
                          ('Marked Absent', 'was Late')])
        self.assertEqual(log[0]['actor_role'], 'admin')
        html = client.get('/admin/logs').get_data(as_text=True)
        self.assertIn('Marked Late', html)

    def test_invalid_changes_are_rejected(self):
        client = self.admin()
        self.post(client, '/admin/attendance', {'register_number': 'R1', 'subject': 'Maths', 'date': TODAY,
                                                'status': 'Maybe'})
        self.post(client, '/admin/attendance', {'register_number': 'R1', 'subject': 'Nope', 'date': TODAY,
                                                'status': 'Present'})
        self.post(client, '/admin/attendance', {'register_number': 'R1', 'subject': 'Maths', 'date': 'bad',
                                                'status': 'Present'})
        self.post(client, '/admin/attendance', {'register_number': 'R9', 'subject': 'Maths', 'date': TODAY,
                                                'status': 'Present'})
        self.assertEqual(self.query('SELECT COUNT(*) FROM attendance')[0][0], 0)
        self.assertEqual(self.query('SELECT COUNT(*) FROM audit_log')[0][0], 0)

    def test_staff_can_only_change_today_for_assigned_students(self):
        self.enroll('R1', 'Maths')
        client, _ = self.login('staff', 'Maths')
        self.post(client, '/staff/attendance', {'register_number': 'R1', 'date': days_ago(1), 'status': 'Present'})
        self.assertIsNone(self.status('R1', date=days_ago(1)))
        self.post(client, '/staff/attendance', {'register_number': 'R2', 'date': TODAY, 'status': 'Present'})
        self.assertIsNone(self.status('R2'))
        self.post(client, '/staff/attendance', {'register_number': 'R1', 'date': TODAY, 'status': 'Present',
                                                'subject': 'Physics'})
        self.assertEqual(self.status('R1')['status'], 'Present')
        self.assertEqual(self.query('SELECT actor FROM audit_log')[0][0], 'Prof (Maths)')
        html = client.get('/staff-dashboard').get_data(as_text=True)
        self.assertIn('Mark Absent', html)
        html = client.get('/staff-dashboard?date=' + days_ago(1)).get_data(as_text=True)
        self.assertIn('read-only', html)
        self.assertNotIn('Mark Absent', html)


class TestReports(WebTestCase):
    def setUp(self):
        super().setUp()
        self.add_subject('Maths')
        self.add_student('R1', 'Arun')
        self.add_student('R2', 'Priya')
        self.dates = [days_ago(n) for n in (4, 3, 2, 1)]
        for i, date in enumerate(self.dates):
            self.mark('R1', 'Maths', date, 'Late' if i == 0 else 'Present')
        self.mark('R2', 'Maths', self.dates[0])
        self.mark('R2', 'Maths', self.dates[3])

    def summary(self, **kwargs):
        conn = self.db()
        try:
            return {r['register_number']: r for r in reports.subject_summary(conn, 'Maths', **kwargs)}
        finally:
            conn.close()

    def test_percentages_and_low_flag(self):
        rows = self.summary()
        self.assertEqual((rows['R1']['held'], rows['R1']['present'], rows['R1']['late'], rows['R1']['absent']),
                         (4, 3, 1, 0))
        self.assertEqual(rows['R1']['percent'], 100.0)
        self.assertEqual((rows['R2']['present'], rows['R2']['absent'], rows['R2']['percent']), (2, 2, 50.0))
        self.assertTrue(rows['R2']['low'])
        self.assertFalse(rows['R1']['low'])

    def test_date_range(self):
        rows = self.summary(start=self.dates[1], end=self.dates[2])
        self.assertEqual(rows['R1']['held'], 2)
        self.assertEqual(rows['R2']['percent'], 0.0)

    def test_general_only_listed_once_used(self):
        conn = self.db()
        try:
            self.assertEqual(common.subjects_for_student(conn, 'R1'), ['Maths'])
        finally:
            conn.close()
        self.mark('R2', common.GENERAL_SUBJECT, TODAY)
        conn = self.db()
        try:
            self.assertEqual(common.subjects_for_student(conn, 'R1'), [common.GENERAL_SUBJECT, 'Maths'])
        finally:
            conn.close()

    def test_report_pages(self):
        html = self.admin().get('/admin/reports').get_data(as_text=True)
        self.assertIn('row-low', html)
        self.assertIn('50.0%', html)
        staff, _ = self.login('staff', 'Maths')
        self.assertIn('100.0%', staff.get('/staff/reports').get_data(as_text=True))

    def test_student_dashboard(self):
        student, _ = self.login('student', 'R2')
        html = student.get('/student-dashboard').get_data(as_text=True)
        self.assertIn('below 75% in: Maths', html)
        self.assertIn('50.0%', html)
        self.assertEqual(html.count('bg-danger">Absent'), 2)
        html = student.get(f'/student-dashboard?start={self.dates[3]}').get_data(as_text=True)
        self.assertIn('100.0%', html)


class TestExports(WebTestCase):
    def setUp(self):
        super().setUp()
        self.add_subject('Maths')
        self.add_subject('Physics')
        self.add_student('R1', '=HYPERLINK("http://evil","x")')
        self.add_student('R2', 'Priya')
        self.mark('R1', 'Maths', days_ago(2))
        self.mark('R1', 'Maths', days_ago(1))
        self.mark('R2', 'Physics', days_ago(1))

    def workbook(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment;', response.headers['Content-Disposition'])
        return load_workbook(io.BytesIO(response.data))

    def test_admin_excel_export(self):
        wb = self.workbook(self.admin().get('/export/xlsx'))
        summary = list(wb['Summary'].iter_rows(values_only=True))
        header_index = summary.index(tuple(reports.SUMMARY_HEADERS))
        data = summary[header_index + 1:]
        self.assertEqual(len(data), 4)          # 2 students x 2 subjects
        names = {row[1] for row in data}
        self.assertIn('=HYPERLINK("http://evil","x")', names)
        self.assertEqual(wb['Summary'].cell(row=header_index + 2, column=2).data_type, 's')
        records = list(wb['Records'].iter_rows(values_only=True))
        statuses = [row[4] for row in records[records.index(tuple(reports.DETAIL_HEADERS)) + 1:]]
        # Maths: R1 present twice, R2 absent twice; Physics: R2 present, R1 absent
        self.assertEqual(sorted(statuses), ['Absent'] * 3 + ['Present'] * 3)

    def test_pdf_export(self):
        response = self.admin().get('/export/pdf?subject=Maths')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data.startswith(b'%PDF'))
        self.assertIn(b'/Type /Page', response.data)

    def test_staff_and_student_exports_are_scoped(self):
        staff, _ = self.login('staff', 'Maths')
        wb = self.workbook(staff.get('/export/xlsx?subject=Physics'))
        subjects = {row[2] for row in wb['Summary'].iter_rows(min_row=5, values_only=True) if row[2]}
        self.assertEqual(subjects, {'Maths'})
        student, _ = self.login('student', 'R2')
        wb = self.workbook(student.get('/export/xlsx'))
        regs = {row[0] for row in wb['Summary'].iter_rows(min_row=5, values_only=True) if row[0]}
        self.assertEqual(regs, {'R2'})


class TestBulkImport(WebTestCase):
    def upload(self, client, text):
        return self.post(client, '/admin/students/import',
                         {'csv_file': (io.BytesIO(text.encode('utf-8-sig')), 'students.csv')},
                         content_type='multipart/form-data')

    def test_import_adds_updates_and_reports_errors(self):
        self.add_subject('Maths')
        self.add_student('R9', 'Old Name')
        csv_text = ('Reg No,Student Name,Password,Subjects\r\n'
                    'R1,Arun,secret11,Maths\r\n'
                    'R2,Priya,secret22,Maths;Chemistry\r\n'
                    'R9,New Name,,\r\n'
                    'R_3,Bad,secret33,\r\n'
                    'R4,NoPassword,,\r\n'
                    'R1,Duplicate,secret11,\r\n'
                    ',,,\r\n')
        client = self.admin()
        self.upload(client, csv_text)
        students = {r['register_number']: r for r in self.query('SELECT * FROM students')}
        self.assertEqual(set(students), {'R1', 'R2', 'R9'})
        self.assertEqual(students['R9']['name'], 'New Name')
        self.assertTrue(common.verify_password(students['R9']['password'], PASSWORD))
        self.assertTrue(common.verify_password(students['R1']['password'], 'secret11'))
        self.assertEqual(sorted(r[0] for r in self.query("SELECT register_number FROM enrollments")), ['R1', 'R2'])
        html = client.get('/admin/students').get_data(as_text=True)
        self.assertIn('2 added, 1 updated, 2 subject assignment(s)', html)
        self.assertIn('unknown subject', html)
        self.assertIn("Line 5: Register number cannot contain", html)
        self.assertIn('Line 6: a password is required', html)
        self.assertIn('Line 7: register number appears more than once', html)

    def test_import_needs_header(self):
        client = self.admin()
        self.upload(client, 'R1,Arun,secret11\r\n')
        self.assertEqual(self.query('SELECT COUNT(*) FROM students')[0][0], 0)
        self.assertIn('header', client.get('/admin/students').get_data(as_text=True))

    def test_sample_csv_imports_cleanly(self):
        self.add_subject('Maths')
        self.add_subject('Physics')
        client = self.admin()
        sample = client.get('/admin/students/sample.csv').get_data(as_text=True)
        self.upload(client, sample)
        self.assertEqual(self.query('SELECT COUNT(*) FROM students')[0][0], 2)
        self.assertEqual(self.query('SELECT COUNT(*) FROM enrollments')[0][0], 2)


class TestBackgroundProcesses(WebTestCase):
    def test_camera_tools_cannot_overlap(self):
        client = self.admin()
        self.assertEqual(self.post_json(client, '/run-register').get_json()['status'], 'success')
        self.assertEqual(self.post_json(client, '/run-register').get_json()['status'], 'error')
        result = self.post_json(client, '/run-attendance', {'subject': 'General'}).get_json()
        self.assertIn('registration window', result['message'])
        self.assertTrue(client.get('/api/process-status').get_json()['register'])
        self.started[0].running = False
        self.assertEqual(self.post_json(client, '/run-attendance', {'subject': 'General'}).get_json()['status'],
                         'success')
        status = client.get('/api/process-status').get_json()
        self.assertEqual((status['scanner'], status['scanner_subject']), (True, 'General'))

    def test_scanner_arguments(self):
        self.add_subject('-Maths')
        staff, _ = self.login('staff', '-Maths')
        result = self.post_json(staff, '/run-attendance', {'subject': 'Hacked', 'duration': '45',
                                                           'late_after': '7.5'}).get_json()
        self.assertEqual(result['status'], 'success')
        args = self.started[-1].args
        self.assertEqual(args[1:], ['attendance_taker.py', '--duration=45', '--late-after=7.5',
                                    '--started-by=Prof (-Maths)', '--', '-Maths'])
        self.started[-1].running = False
        admin = self.admin()
        self.assertEqual(self.post_json(admin, '/run-attendance', {'subject': 'Unknown'}).get_json()['status'], 'error')
        self.post_json(admin, '/run-attendance', {'subject': 'General', 'duration': 'abc', 'late_after': -5})
        self.assertEqual(self.started[-1].args[2:4], ['--duration=0', '--late-after=0'])

    def test_extraction_and_evaluation_run_once_at_a_time(self):
        client = self.admin()
        self.assertEqual(self.post_json(client, '/run-extract').get_json()['status'], 'success')
        self.assertEqual(self.post_json(client, '/run-extract').get_json()['status'], 'error')
        self.assertEqual(self.post_json(client, '/run-evaluation').get_json()['status'], 'success')
        self.assertEqual(self.post_json(client, '/run-evaluation').get_json()['status'], 'error')

    def test_missing_models_are_reported(self):
        with mock.patch.object(common, 'missing_model_files', return_value=[common.SHAPE_PREDICTOR_FILE]):
            result = self.post_json(self.admin(), '/run-register').get_json()
        self.assertEqual(result['status'], 'error')
        self.assertIn('download_models.py', result['message'])
        self.assertEqual(self.started, [])


class TestSessions(WebTestCase):
    def test_recent_sessions_newest_first_with_day_counts(self):
        self.add_subject('Maths')
        self.add_student('R1', 'Arun')
        with self.db() as conn:
            for day in (days_ago(1), days_ago(3), days_ago(2)):     # inserted out of order
                conn.execute('INSERT INTO sessions (subject, date, started_at, duration_min, late_after_min) '
                             'VALUES (?, ?, ?, 60, 10)', ('Maths', day, day + ' 09:00:00'))
        self.mark('R1', 'Maths', days_ago(2), 'Late')
        conn = self.db()
        try:
            rows = webapp.recent_sessions(conn, 'Maths')
        finally:
            conn.close()
        self.assertEqual([r['date'] for r in rows], [days_ago(1), days_ago(2), days_ago(3)])
        self.assertEqual([(r['present_count'], r['late_count']) for r in rows], [(0, 0), (0, 1), (0, 0)])
        staff, _ = self.login('staff', 'Maths')
        self.assertIn('1 late', staff.get('/staff-dashboard').get_data(as_text=True))


class TestAccuracyPage(WebTestCase):
    def test_shows_saved_evaluation(self):
        result = {'evaluated_at': '2026-10-03 10:00:00', 'configured_threshold': common.MATCH_THRESHOLD,
                  'dataset': {'students': 2, 'photos': 20, 'faces_detected': 20, 'genuine_trials': 20,
                              'impostor_trials': 20},
                  'thresholds': [{'threshold': common.MATCH_THRESHOLD, 'genuine_trials': 20, 'correct_matches': 19,
                                  'false_rejects': 1, 'wrong_person': 0, 'impostor_trials': 20,
                                  'false_accepts': 0, 'true_rejects': 20, 'tar': 95.0, 'frr': 5.0,
                                  'wrong_rate': 0.0, 'far': 0.0, 'accuracy': 97.5}],
                  'recommended_threshold': common.MATCH_THRESHOLD,
                  'latency_ms': {'color_conversion': 0.5, 'detection': 80.0, 'landmarks': 2.0, 'descriptor': 20.0,
                                 'matching': 0.01, 'per_frame_total': 102.5, 'fps_estimate': 9.8,
                                 'sqlite_write': 3.0},
                  'notes': ['note one']}
        with open(common.EVALUATION_JSON, 'w', encoding='utf-8') as f:
            json.dump(result, f)
        client = self.admin()
        html = client.get('/admin/accuracy').get_data(as_text=True)
        self.assertIn('97.5%', html)
        self.assertIn('note one', html)
        self.assertIn('97.5%', client.get('/admin-dashboard').get_data(as_text=True))


if __name__ == '__main__':
    unittest.main()
